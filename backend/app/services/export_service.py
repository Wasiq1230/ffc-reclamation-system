import csv
from io import StringIO, BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def export_jobs_to_csv(jobs_data: list[dict]) -> str:
    """Exports list of jobs as CSV text."""
    output = StringIO()
    headers = [
        "Tracking ID", "Registration Date", "Source", "Plant", "Department",
        "External Customer", "Notification No", "Valve Description",
        "Valve Serial No", "Valve Tag No", "Size", "Rating", "Quantity",
        "Status", "Technician", "Created At"
    ]
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for j in jobs_data:
        writer.writerow({
            "Tracking ID": j.get("tracking_id", ""),
            "Registration Date": j.get("registration_date", ""),
            "Source": j.get("source", ""),
            "Plant": j.get("plant_name", ""),
            "Department": j.get("department_name", ""),
            "External Customer": j.get("customer_name", ""),
            "Notification No": j.get("notification_no", ""),
            "Valve Description": j.get("valve_description", ""),
            "Valve Serial No": j.get("valve_serial_no", ""),
            "Valve Tag No": j.get("valve_tag_no", ""),
            "Size": j.get("size", ""),
            "Rating": j.get("pressure_rating", ""),
            "Quantity": j.get("valve_quantity", 1),
            "Status": j.get("status", ""),
            "Technician": j.get("technician_name", ""),
            "Created At": j.get("created_at", "")
        })
    return output.getvalue()

def export_jobs_to_excel(jobs_data: list[dict], title: str = "FFC Valve Reclamation Register") -> bytes:
    """Exports list of jobs as formatted Excel workbook (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Reclamation Records"
    
    # Title row
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    headers = [
        "Tracking ID", "Reg Date", "Source", "Plant", "Department",
        "Customer", "Notification", "Valve Description", "Tag / Serial",
        "Size & Rating", "Qty", "Status", "Technician", "Registered Date"
    ]
    
    ws.append([]) # Row 2 empty
    ws.append(headers) # Row 3 headers
    
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 24

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for row_idx, j in enumerate(jobs_data, start=4):
        tag_serial = f"{j.get('valve_tag_no', '') or ''} / {j.get('valve_serial_no', '') or ''}".strip(" /")
        size_rating = f"{j.get('size', '') or ''} {j.get('pressure_rating', '') or ''}".strip()
        row = [
            j.get("tracking_id", ""),
            str(j.get("date_received", "")),
            j.get("source", ""),
            j.get("plant_name", ""),
            j.get("department_name", ""),
            j.get("customer_name", ""),
            j.get("notification_no", ""),
            j.get("valve_description", ""),
            tag_serial,
            size_rating,
            j.get("valve_quantity", 1),
            j.get("status", ""),
            j.get("technician_name", ""),
            str(j.get("registration_date", ""))[:10]
        ]
        ws.append(row)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=9)
            if col_num in [1, 2, 3, 4, 5, 7, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="center")

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
